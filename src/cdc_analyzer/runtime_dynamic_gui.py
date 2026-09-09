from __future__ import annotations

from pathlib import Path
from tempfile import TemporaryDirectory

import numpy as np
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage

from .dynamic_analysis import CURRENT, LOAD, TIME, VELOCITY, ResponseConfig, ResponseStandard
from .dynamic_export import export_hysteresis_xlsx, export_response_xlsx
from .dynamic_gui import DynamicPagesController as _BaseDynamicPagesController
from .response_v074 import analyze_response_time_v074
from .response_v075 import analyze_response_time_v075
from .response_v080 import analyze_response_time_v080, default_target_speeds, parse_target_speeds


class _AnalyzerModuleProxy:
    """Compatibility proxy for historical controller analyzer redirection."""
    def __getattr__(self, name):
        return globals()[name]

    def __setattr__(self, name, value):
        globals()[name] = value


_v074_module = _AnalyzerModuleProxy()
_dynamic_gui_v074_module = _v074_module
_v074_module.analyze_response_time_v074 = analyze_response_time_v075


# Consolidated behavior layer 074
class _Layer074(_BaseDynamicPagesController):
    """V0.7.4 response-page upgrade.

    Hysteresis behavior remains inherited from the released controller. Only
    response event selection, target-speed evaluation display, plot density and
    readability are changed here.
    """

    def _new_table(self):
        table = super()._new_table()
        font = table.font()
        font.setPointSize(max(10, font.pointSize()))
        table.setFont(font)
        header_font = table.horizontalHeader().font()
        header_font.setPointSize(max(10, header_font.pointSize()))
        header_font.setBold(True)
        table.horizontalHeader().setFont(header_font)
        table.verticalHeader().setDefaultSectionSize(30)
        table.setSelectionBehavior(
            self.QtWidgets.QAbstractItemView.SelectionBehavior.SelectRows
        )
        return table

    def _build_response_page(self):
        super()._build_response_page()
        self.response_event_combo.setMinimumWidth(360)
        self.response_table.currentCellChanged.connect(
            self._select_response_event_from_table
        )

    def _select_response_event_from_table(
        self,
        current_row: int,
        _current_column: int,
        _previous_row: int,
        _previous_column: int,
    ):
        if (
            self.response_result is None
            or current_row < 0
            or current_row >= len(self.response_result.events)
        ):
            return
        event_id = int(self.response_result.events.iloc[current_row]["Event ID"])
        index = self.response_event_combo.findData(event_id)
        if index >= 0 and index != self.response_event_combo.currentIndex():
            self.response_event_combo.setCurrentIndex(index)

    def analyze_response(self):
        if self.response_dataset is None:
            self.QtWidgets.QMessageBox.information(
                self.window,
                self._text("响应时间", "Response Time"),
                self._text(
                    "请先打开响应时间原始数据。",
                    "Open response-time raw data first.",
                ),
            )
            return
        try:
            standard = ResponseStandard(self.response_standard.currentData())
            limit = self.response_t90_limit.value() or None
            config = ResponseConfig(
                standard=standard,
                trigger_fraction=self.response_trigger.value() / 100.0,
                end_average_fraction=self.response_end_fraction.value() / 100.0,
                t90_limit_ms=limit,
            )
            self.response_result = analyze_response_time_v074(
                self.response_dataset,
                config,
                target_speed_tolerance=0.10,
            )
            self._fill_table(self.response_table, self.response_result.events)

            self.response_event_combo.blockSignals(True)
            self.response_event_combo.clear()
            for _, row in self.response_result.events.iterrows():
                stage = str(row.get("Stage", ""))
                direction = (
                    self._text("复原(+)", "Rebound (+)")
                    if row["Direction"] == "Rebound"
                    else self._text("压缩(-)", "Compression (-)")
                )
                target = float(row["Target Velocity m/s"])
                current_transition = str(row.get("Current Transition", ""))
                text = (
                    f"{stage} | {direction} | {current_transition} | "
                    f"{self._text('目标速度', 'Target')} {target:+.4g} m/s"
                )
                self.response_event_combo.addItem(text, int(row["Event ID"]))
            self.response_event_combo.blockSignals(False)

            if self.response_event_combo.count():
                self.response_event_combo.setCurrentIndex(0)
            self.refresh_response_plot()

            settings = self.response_result.settings
            detected = int(settings.get("Detected Current Events", len(self.response_result.events)))
            accepted = int(settings.get("Accepted Target-Speed Events", len(self.response_result.events)))
            rejected = int(settings.get("Rejected Non-target Events", 0))
            invalid = int(settings.get("Rejected Invalid Events", 0))
            sample_rate = float(self.response_result.events["Sample Rate Hz"].iloc[0])
            self.response_status.setText(
                self._text(
                    (
                        f"目标速度响应：检测 {detected} 个电流切换，保留 {accepted} 个，"
                        f"排除非目标速度 {rejected} 个、无效 {invalid} 个；"
                        f"采样率约 {sample_rate:.2f} Hz"
                    ),
                    (
                        f"Target-speed response: {detected} current steps detected, "
                        f"{accepted} retained, {rejected} off-target and {invalid} invalid "
                        f"event(s) excluded; sample rate ≈ {sample_rate:.2f} Hz"
                    ),
                )
            )
        except Exception as exc:
            self.QtWidgets.QMessageBox.critical(
                self.window,
                self._text("分析错误", "Analysis error"),
                str(exc),
            )

    def _axis_style(self, plot, left: str, units: str | None = None):
        plot.showGrid(x=True, y=False, alpha=0.12)
        label_style = {"font-size": "11pt"}
        plot.setLabel("left", left, units=units, **label_style)
        plot.setLabel(
            "bottom",
            self._text("时间", "Time"),
            units="ms",
            **label_style,
        )
        plot.setClipToView(True)
        plot.setDownsampling(auto=True, mode="peak")
        background = getattr(self.window, "_plot_background_color", "#ffffff")
        foreground = getattr(self.window, "_plot_foreground_color", "#202020")
        plot.getAxis("left").setTextPen(foreground)
        plot.getAxis("bottom").setTextPen(foreground)
        tick_font = self.QtWidgets.QApplication.font()
        tick_font.setPointSize(max(10, tick_font.pointSize()))
        plot.getAxis("left").setStyle(tickFont=tick_font)
        plot.getAxis("bottom").setStyle(tickFont=tick_font)
        self.response_plot_area.setBackground(background)
        self.hysteresis_plot_area.setBackground(background)

    def _marker_pen(self):
        foreground = getattr(self.window, "_plot_foreground_color", "#202020")
        return self.pg.mkPen(
            foreground,
            width=1.1,
            style=self.QtCore.Qt.PenStyle.DashLine,
        )

    def _signal_pen(self):
        return self.pg.mkPen("#1565c0", width=2.2)

    def _add_plot_text(self, plot, text: str, x: float, y: float, *, anchor=(0, 0.5), bold=False):
        foreground = getattr(self.window, "_plot_foreground_color", "#202020")
        item = self.pg.TextItem(text=text, color=foreground, anchor=anchor)
        font = self.QtWidgets.QApplication.font()
        font.setPointSize(max(10, font.pointSize()))
        font.setBold(bool(bold))
        item.setFont(font)
        item.setPos(float(x), float(y))
        plot.addItem(item)
        return item

    def _sync_response_table_selection(self, event_id: int):
        if self.response_result is None or self.response_result.events.empty:
            return
        matches = np.flatnonzero(
            self.response_result.events["Event ID"].to_numpy(int) == int(event_id)
        )
        if not len(matches):
            return
        row_index = int(matches[0])
        if self.response_table.currentRow() == row_index:
            return
        self.response_table.blockSignals(True)
        self.response_table.selectRow(row_index)
        self.response_table.blockSignals(False)

    def refresh_response_plot(self):
        self.response_plot_area.clear()
        if self.response_result is None or self.response_result.events.empty:
            return

        event_id = self.response_event_combo.currentData()
        if event_id is None:
            event_id = int(self.response_result.events["Event ID"].iloc[0])
        row = self.response_result.events[
            self.response_result.events["Event ID"] == event_id
        ].iloc[0]
        self._sync_response_table_selection(int(event_id))

        start = float(row.get("Display Start s", row["Segment Start s"]))
        end = float(row.get("Display End s", row["Segment End s"]))
        data = self.response_result.processed[
            self.response_result.processed[TIME].between(start, end)
        ]
        if data.empty:
            return

        t_ms = data[TIME].to_numpy(float) * 1000.0
        t0_ms = float(row["t0 s"]) * 1000.0
        x_left = float(t_ms[0])
        x_right = float(t_ms[-1])
        x_label = x_left + 0.02 * max(x_right - x_left, 1e-6)
        marker_pen = self._marker_pen()
        signal_pen = self._signal_pen()

        current_plot = self.response_plot_area.addPlot(row=0, col=0)
        self._axis_style(current_plot, self._text("阀电流", "Valve current"), "A")
        current_plot.plot(
            t_ms,
            data[CURRENT].to_numpy(float),
            pen=signal_pen,
        )
        current_plot.addLine(x=t0_ms, pen=marker_pen)
        current_levels = (
            ("10%", float(row["Trigger Current A"])),
            ("100%", float(row["Current 100% A"])),
        )
        for label, value in current_levels:
            current_plot.addLine(y=value, pen=marker_pen)
            self._add_plot_text(current_plot, label, x_label, value)
        current_plot.setTitle(
            self._text(
                f"电流｜{row.get('Stage', '')}｜{row['Direction']}",
                f"Current | {row.get('Stage', '')} | {row['Direction']}",
            )
        )

        force_plot = self.response_plot_area.addPlot(row=1, col=0)
        force_plot.setXLink(current_plot)
        self._axis_style(force_plot, self._text("阻尼力", "Damping force"), "kN")
        force_kn = data[LOAD].to_numpy(float) / 1000.0
        force_plot.plot(t_ms, force_kn, pen=signal_pen)
        force_plot.addLine(x=t0_ms, pen=marker_pen)

        force_levels = (
            ("1%", float(row["F1 N"]) / 1000.0),
            ("63%", float(row["F63 N"]) / 1000.0),
            ("90%", float(row["F90 N"]) / 1000.0),
            ("100%", float(row["F100 N"]) / 1000.0),
        )
        for label, value in force_levels:
            force_plot.addLine(y=value, pen=marker_pen)
            self._add_plot_text(force_plot, label, x_label, value)

        response_markers = (
            ("t1", float(row["Dead Time t1 ms"])),
            ("t63", float(row["Switch Time t63 ms"])),
            ("t90", float(row["Switch Time t90 ms"])),
        )
        finite_marker_positions: list[tuple[str, float]] = []
        for label, elapsed_ms in response_markers:
            if np.isfinite(elapsed_ms):
                x_value = t0_ms + elapsed_ms
                force_plot.addLine(x=x_value, pen=marker_pen)
                finite_marker_positions.append((label, x_value))
        if len(force_kn):
            y_min = float(np.nanmin(force_kn))
            y_max = float(np.nanmax(force_kn))
            y_span = max(y_max - y_min, 0.1)
            for index, (label, x_value) in enumerate(finite_marker_positions):
                y = y_max - (0.08 + 0.12 * index) * y_span
                self._add_plot_text(force_plot, label, x_value, y, anchor=(0.5, 0.5))

        direction_text = (
            self._text("复原（+）", "Rebound (+)")
            if row["Direction"] == "Rebound"
            else self._text("压缩（-）", "Compression (-)")
        )
        if len(force_kn):
            direction_y = float(np.nanmax(force_kn)) if row["Direction"] == "Rebound" else float(np.nanmin(force_kn))
            self._add_plot_text(
                force_plot,
                direction_text,
                x_left + 0.55 * max(x_right - x_left, 1e-6),
                direction_y,
                anchor=(0.5, 1.0 if row["Direction"] == "Rebound" else 0.0),
                bold=True,
            )

        velocity_plot = self.response_plot_area.addPlot(row=2, col=0)
        velocity_plot.setXLink(current_plot)
        self._axis_style(velocity_plot, self._text("速度", "Velocity"), "m/s")
        velocity_plot.plot(
            t_ms,
            data[VELOCITY].to_numpy(float),
            pen=signal_pen,
        )
        velocity_plot.addLine(x=t0_ms, pen=marker_pen)
        target_velocity = float(row["Target Velocity m/s"])
        velocity_plot.addLine(y=target_velocity, pen=marker_pen)
        self._add_plot_text(
            velocity_plot,
            self._text(
                f"目标 {target_velocity:+.4g} m/s",
                f"Target {target_velocity:+.4g} m/s",
            ),
            x_label,
            target_velocity,
            bold=True,
        )

        current_plot.setXRange(x_left, x_right, padding=0.01)


# Consolidated behavior layer 075
class _Layer075(_Layer074):
    """V0.7.5 dynamic pages.

    Response plots and response results use separate sub-tabs so the three
    synchronized plots can consume the full available page height.
    """

    RESPONSE_HEADERS_ZH = {
        "Event ID": "事件编号",
        "Detected Event ID": "原始事件编号",
        "OEM": "规范",
        "Stage": "切换阶段",
        "Current Transition": "电流切换",
        "Current Start A": "起始电流 / A",
        "Current End A": "终止电流 / A",
        "Current Delta A": "电流变化 / A",
        "Trigger Fraction": "触发比例",
        "Trigger Current A": "I₁₀% / A",
        "Current 100% A": "I₁₀₀% / A",
        "t0 s": "t₀ / s",
        "Displacement at t0 mm": "t0位移 / mm",
        "Velocity at t0 m/s": "t0速度 / m/s",
        "Target Velocity m/s": "目标速度 / m/s",
        "Target Speed Error %": "速度误差 / %",
        "Direction": "方向",
        "Force Change": "载荷变化",
        "F0 N": "F0 / N",
        "F1 N": "F₁% / N",
        "F63 N": "F₆₃% / N",
        "F90 N": "F₉₀% / N",
        "F100 N": "F₁₀₀% / N",
        "Delta F N": "ΔF / N",
        "Dead Time t1 ms": "t₁% / ms",
        "Switch Time t63 ms": "t₆₃% / ms",
        "Switch Time t90 ms": "t₉₀% / ms",
        "Gradient 63 N/s": "梯度63 / N/s",
        "Gradient 90 N/s": "梯度90 / N/s",
        "Sample Rate Hz": "采样率 / Hz",
        "Status": "状态",
        "Issues": "数据提示",
    }

    def _build_response_page(self):
        super()._build_response_page()

        root = self.response_page.layout()
        splitter = self.response_plot_area.parentWidget()
        if splitter is not None:
            root.removeWidget(splitter)
            self.response_plot_area.setParent(None)
            self.response_table.setParent(None)
            splitter.setParent(None)

        self.response_view_tabs = self.QtWidgets.QTabWidget()
        self.response_view_tabs.setDocumentMode(False)

        self.response_graph_page = self.QtWidgets.QWidget()
        graph_layout = self.QtWidgets.QVBoxLayout(self.response_graph_page)
        graph_layout.setContentsMargins(0, 0, 0, 0)
        graph_layout.setSpacing(4)
        graph_layout.addWidget(self.response_plot_area, 1)

        self.response_data_page = self.QtWidgets.QWidget()
        data_layout = self.QtWidgets.QVBoxLayout(self.response_data_page)
        data_layout.setContentsMargins(0, 0, 0, 0)
        data_layout.setSpacing(6)

        data_tools = self.QtWidgets.QHBoxLayout()
        self.response_data_hint = self.QtWidgets.QLabel()
        self.response_data_hint.setWordWrap(True)
        data_tools.addWidget(self.response_data_hint, 1)
        self.response_show_plot_button = self.QtWidgets.QPushButton()
        self.response_show_plot_button.clicked.connect(self._show_selected_response_plot)
        data_tools.addWidget(self.response_show_plot_button)
        data_layout.addLayout(data_tools)
        data_layout.addWidget(self.response_table, 1)

        self.response_view_tabs.addTab(self.response_graph_page, "")
        self.response_view_tabs.addTab(self.response_data_page, "")
        root.addWidget(self.response_view_tabs, 1)

        self.response_table.cellDoubleClicked.connect(
            lambda _row, _column: self._show_selected_response_plot()
        )

    def _show_selected_response_plot(self):
        if hasattr(self, "response_view_tabs"):
            self.response_view_tabs.setCurrentWidget(self.response_graph_page)
        self.refresh_response_plot()

    def _fill_table(self, table, frame):
        super()._fill_table(table, frame)
        if (
            hasattr(self, "response_table")
            and table is self.response_table
            and self.window.language == "zh_CN"
        ):
            labels = [
                self.RESPONSE_HEADERS_ZH.get(str(column), str(column))
                for column in frame.columns
            ]
            table.setHorizontalHeaderLabels(labels)

    def apply_language(self, language: str):
        super().apply_language(language)
        if hasattr(self, "response_view_tabs"):
            self.response_view_tabs.setTabText(
                self.response_view_tabs.indexOf(self.response_graph_page),
                self._text("图形分析", "Plot Analysis"),
            )
            self.response_view_tabs.setTabText(
                self.response_view_tabs.indexOf(self.response_data_page),
                self._text("结果数据", "Result Data"),
            )
            self.response_data_hint.setText(
                self._text(
                    "选择任一结果行后，可点击“查看图形”或双击该行查看对应响应阶段。",
                    "Select a result row, then click View Plot or double-click the row to inspect that response event.",
                )
            )
            self.response_show_plot_button.setText(
                self._text("查看图形", "View Plot")
            )

        if hasattr(self, "response_table") and self.response_result is not None:
            self._fill_table(self.response_table, self.response_result.events)


# Consolidated behavior layer 076
class _Layer076(_Layer075):
    """V0.7.6 dynamic pages.

    Response and hysteresis now consume the shared data source loaded from the
    main left-side Data File panel. Module-local Open/Export buttons remain as
    compatibility objects but are hidden from the user-facing layout.
    """

    def _build_response_page(self):
        super()._build_response_page()
        self.response_open_button.hide()
        self.response_export_button.hide()

    def _build_hysteresis_page(self):
        super()._build_hysteresis_page()
        self.hysteresis_open_button.hide()
        self.hysteresis_export_button.hide()

    def set_shared_dataset(self, dataset, path: str | Path | None = None) -> None:
        source_path = Path(path) if path is not None else Path(dataset.source_path)
        self.response_dataset = dataset
        self.response_path = source_path
        self.hysteresis_dataset = dataset
        self.hysteresis_path = source_path
        self.response_result = None
        self.hysteresis_result = None
        self._update_shared_source_labels()
        self.response_status.setText(
            self._text("已使用左侧公共数据源；点击“分析响应”开始计算。", "Shared data source loaded; click Analyze Response.")
        )
        self.hysteresis_status.setText(
            self._text("已使用左侧公共数据源；点击“分析迟滞”开始计算。", "Shared data source loaded; click Analyze Hysteresis.")
        )

    def _update_shared_source_labels(self) -> None:
        if self.response_path is not None:
            self.response_file_label.setText(
                self._text(f"公共数据源：{self.response_path}", f"Shared source: {self.response_path}")
            )
        else:
            self.response_file_label.setText(
                self._text("公共数据源：未加载", "Shared source: not loaded")
            )
        if self.hysteresis_path is not None:
            self.hysteresis_file_label.setText(
                self._text(f"公共数据源：{self.hysteresis_path}", f"Shared source: {self.hysteresis_path}")
            )
        else:
            self.hysteresis_file_label.setText(
                self._text("公共数据源：未加载", "Shared source: not loaded")
            )

    def apply_language(self, language: str):
        super().apply_language(language)
        if hasattr(self, "response_open_button"):
            self.response_open_button.hide()
            self.response_export_button.hide()
            self.hysteresis_open_button.hide()
            self.hysteresis_export_button.hide()
            self._update_shared_source_labels()

    def _export_plot_widget_png(self, plot_widget, path: str | Path, scale: float = 3.0) -> Path:
        import pyqtgraph.exporters

        out = Path(path).with_suffix(".png")
        out.parent.mkdir(parents=True, exist_ok=True)
        exporter = pyqtgraph.exporters.ImageExporter(plot_widget.scene())
        params = exporter.parameters()
        base_width = max(int(plot_widget.width()), 800)
        params["width"] = max(2400, int(base_width * float(scale)))
        exporter.export(str(out))
        return out

    def export_response_png(self) -> None:
        if self.response_result is None:
            self.QtWidgets.QMessageBox.information(
                self.window,
                self._text("导出 PNG", "Export PNG"),
                self._text("请先完成响应时间分析。", "Analyze response time first."),
            )
            return
        default = (
            self.response_path.stem + "_response.png"
            if self.response_path is not None
            else "response.png"
        )
        path, _ = self.QtWidgets.QFileDialog.getSaveFileName(
            self.window,
            self._text("导出响应时间图片", "Export response-time plot"),
            default,
            "PNG (*.png)",
        )
        if not path:
            return
        try:
            if hasattr(self, "response_view_tabs"):
                self.response_view_tabs.setCurrentWidget(self.response_graph_page)
            self.refresh_response_plot()
            self.QtWidgets.QApplication.processEvents()
            out = self._export_plot_widget_png(self.response_plot_area, path, 3.0)
            self.window.statusBar().showMessage(
                self._text(f"已导出高清 PNG：{out}", f"High-resolution PNG exported: {out}")
            )
        except Exception as exc:
            self.QtWidgets.QMessageBox.critical(
                self.window,
                self._text("导出错误", "Export error"),
                str(exc),
            )

    def export_hysteresis_png(self) -> None:
        if self.hysteresis_result is None:
            self.QtWidgets.QMessageBox.information(
                self.window,
                self._text("导出 PNG", "Export PNG"),
                self._text("请先完成迟滞分析。", "Analyze hysteresis first."),
            )
            return
        default = (
            self.hysteresis_path.stem + "_hysteresis.png"
            if self.hysteresis_path is not None
            else "hysteresis.png"
        )
        path, _ = self.QtWidgets.QFileDialog.getSaveFileName(
            self.window,
            self._text("导出迟滞图片", "Export hysteresis plot"),
            default,
            "PNG (*.png)",
        )
        if not path:
            return
        try:
            self.refresh_hysteresis_plot()
            self.QtWidgets.QApplication.processEvents()
            out = self._export_plot_widget_png(self.hysteresis_plot_area, path, 3.0)
            self.window.statusBar().showMessage(
                self._text(f"已导出高清 PNG：{out}", f"High-resolution PNG exported: {out}")
            )
        except Exception as exc:
            self.QtWidgets.QMessageBox.critical(
                self.window,
                self._text("导出错误", "Export error"),
                str(exc),
            )

    def _append_response_plots(self, workbook_path: Path) -> None:
        if self.response_result is None or self.response_result.events.empty:
            return
        original_index = self.response_event_combo.currentIndex()
        workbook = load_workbook(workbook_path)
        if "Response Plots" in workbook.sheetnames:
            del workbook["Response Plots"]
        sheet = workbook.create_sheet("Response Plots")
        sheet.sheet_view.showGridLines = False
        sheet.column_dimensions["A"].width = 22
        sheet.column_dimensions["B"].width = 24

        with TemporaryDirectory(prefix="damper_response_plots_") as tmp:
            tmp_dir = Path(tmp)
            anchor_row = 1
            for index, row in self.response_result.events.reset_index(drop=True).iterrows():
                event_id = int(row["Event ID"])
                combo_index = self.response_event_combo.findData(event_id)
                if combo_index >= 0:
                    self.response_event_combo.setCurrentIndex(combo_index)
                self.refresh_response_plot()
                self.QtWidgets.QApplication.processEvents()
                image_path = tmp_dir / f"event_{event_id:03d}.png"
                self._export_plot_widget_png(self.response_plot_area, image_path, 2.5)

                stage = str(row.get("Stage", ""))
                direction = str(row.get("Direction", ""))
                current_transition = str(row.get("Current Transition", ""))
                target = row.get("Target Velocity m/s", "")
                sheet.cell(anchor_row, 1).value = (
                    f"Event {event_id} | {stage} | {direction} | {current_transition} | Target {target} m/s"
                )
                image = XLImage(str(image_path))
                if image.width:
                    ratio = min(1.0, 1180.0 / float(image.width))
                    image.width = int(image.width * ratio)
                    image.height = int(image.height * ratio)
                sheet.add_image(image, f"A{anchor_row + 1}")
                rows_used = max(32, int(image.height / 20) + 5)
                anchor_row += rows_used

        if original_index >= 0 and original_index < self.response_event_combo.count():
            self.response_event_combo.setCurrentIndex(original_index)
            self.refresh_response_plot()
        workbook.save(workbook_path)

    def _append_hysteresis_plot(self, workbook_path: Path) -> None:
        if self.hysteresis_result is None:
            return
        workbook = load_workbook(workbook_path)
        if "Hysteresis Plot" in workbook.sheetnames:
            del workbook["Hysteresis Plot"]
        sheet = workbook.create_sheet("Hysteresis Plot")
        sheet.sheet_view.showGridLines = False
        with TemporaryDirectory(prefix="damper_hysteresis_plot_") as tmp:
            image_path = Path(tmp) / "hysteresis.png"
            self.refresh_hysteresis_plot()
            self.QtWidgets.QApplication.processEvents()
            self._export_plot_widget_png(self.hysteresis_plot_area, image_path, 2.5)
            image = XLImage(str(image_path))
            if image.width:
                ratio = min(1.0, 1180.0 / float(image.width))
                image.width = int(image.width * ratio)
                image.height = int(image.height * ratio)
            sheet.add_image(image, "A1")
        workbook.save(workbook_path)

    def export_response(self):
        if self.response_result is None:
            self.QtWidgets.QMessageBox.information(
                self.window,
                self._text("导出 Excel", "Export Excel"),
                self._text("请先完成响应时间分析。", "Analyze response time first."),
            )
            return
        default = (
            self.response_path.stem + "_response.xlsx"
            if self.response_path is not None
            else "response.xlsx"
        )
        path, _ = self.QtWidgets.QFileDialog.getSaveFileName(
            self.window,
            self._text("导出响应时间结果", "Export response-time result"),
            default,
            "Excel (*.xlsx)",
        )
        if not path:
            return
        try:
            out = export_response_xlsx(self.response_result, path)
            self._append_response_plots(out)
            self.window.statusBar().showMessage(
                self._text(f"已导出响应 Excel（全部阶段与图形）：{out}", f"Response Excel exported with all event plots: {out}")
            )
        except Exception as exc:
            self.QtWidgets.QMessageBox.critical(
                self.window,
                self._text("导出错误", "Export error"),
                str(exc),
            )

    def export_hysteresis(self):
        if self.hysteresis_result is None:
            self.QtWidgets.QMessageBox.information(
                self.window,
                self._text("导出 Excel", "Export Excel"),
                self._text("请先完成迟滞分析。", "Analyze hysteresis first."),
            )
            return
        default = (
            self.hysteresis_path.stem + "_hysteresis.xlsx"
            if self.hysteresis_path is not None
            else "hysteresis.xlsx"
        )
        path, _ = self.QtWidgets.QFileDialog.getSaveFileName(
            self.window,
            self._text("导出迟滞结果", "Export hysteresis result"),
            default,
            "Excel (*.xlsx)",
        )
        if not path:
            return
        try:
            out = export_hysteresis_xlsx(self.hysteresis_result, path)
            self._append_hysteresis_plot(out)
            self.window.statusBar().showMessage(
                self._text(f"已导出迟滞 Excel（含图形）：{out}", f"Hysteresis Excel exported with plot: {out}")
            )
        except Exception as exc:
            self.QtWidgets.QMessageBox.critical(
                self.window,
                self._text("导出错误", "Export error"),
                str(exc),
            )


# Consolidated behavior layer 077
class _Layer077(_Layer076):
    """V0.7.7 response-plot presentation refinements.

    The response calculation itself is unchanged. This layer only adjusts the
    detail-plot presentation requested during visual validation:
    - response detail time axis is shown directly in seconds;
    - reference dashes use a wider dash/gap pattern;
    - the force-direction annotation is removed from the force plot;
    - the target-velocity line/text is removed from the velocity plot.
    """

    def _axis_style(self, plot, left: str, units: str | None = None):
        plot.showGrid(x=True, y=False, alpha=0.12)
        label_style = {"font-size": "11pt"}
        plot.setLabel("left", left, units=units, **label_style)
        # Use seconds directly. Passing millisecond-valued X data together with
        # units="ms" allowed pyqtgraph SI-prefix scaling to display "kms".
        plot.setLabel(
            "bottom",
            self._text("时间", "Time"),
            units="s",
            **label_style,
        )
        plot.setClipToView(True)
        plot.setDownsampling(auto=True, mode="peak")
        background = getattr(self.window, "_plot_background_color", "#ffffff")
        foreground = getattr(self.window, "_plot_foreground_color", "#202020")
        plot.getAxis("left").setTextPen(foreground)
        plot.getAxis("bottom").setTextPen(foreground)
        tick_font = self.QtWidgets.QApplication.font()
        tick_font.setPointSize(max(10, tick_font.pointSize()))
        plot.getAxis("left").setStyle(tickFont=tick_font)
        plot.getAxis("bottom").setStyle(tickFont=tick_font)
        self.response_plot_area.setBackground(background)
        self.hysteresis_plot_area.setBackground(background)

    def _marker_pen(self):
        foreground = getattr(self.window, "_plot_foreground_color", "#202020")
        pen = self.pg.mkPen(foreground, width=1.0)
        pen.setStyle(self.QtCore.Qt.PenStyle.CustomDashLine)
        # Qt dash units are multiples of pen width. A 12/8 pattern leaves
        # noticeably larger gaps than the default DashLine pattern while
        # keeping the reference lines light.
        pen.setDashPattern([12.0, 8.0])
        return pen

    def refresh_response_plot(self):
        self.response_plot_area.clear()
        if self.response_result is None or self.response_result.events.empty:
            return

        event_id = self.response_event_combo.currentData()
        if event_id is None:
            event_id = int(self.response_result.events["Event ID"].iloc[0])
        row = self.response_result.events[
            self.response_result.events["Event ID"] == event_id
        ].iloc[0]
        self._sync_response_table_selection(int(event_id))

        start = float(row.get("Display Start s", row["Segment Start s"]))
        end = float(row.get("Display End s", row["Segment End s"]))
        data = self.response_result.processed[
            self.response_result.processed[TIME].between(start, end)
        ]
        if data.empty:
            return

        # Plot the measured running time directly in seconds so the engineering
        # reading is 1.100, 1.110, ... s instead of an SI-prefixed "kms" axis.
        t_s = data[TIME].to_numpy(float)
        t0_s = float(row["t0 s"])
        x_left = float(t_s[0])
        x_right = float(t_s[-1])
        x_label = x_left + 0.02 * max(x_right - x_left, 1e-9)
        marker_pen = self._marker_pen()
        signal_pen = self._signal_pen()

        current_plot = self.response_plot_area.addPlot(row=0, col=0)
        self._axis_style(current_plot, self._text("阀电流", "Valve current"), "A")
        current_plot.plot(t_s, data[CURRENT].to_numpy(float), pen=signal_pen)
        current_plot.addLine(x=t0_s, pen=marker_pen)
        current_levels = (
            ("I10%", float(row["Trigger Current A"])),
            ("I100%", float(row["Current 100% A"])),
        )
        for label, value in current_levels:
            current_plot.addLine(y=value, pen=marker_pen)
            self._add_plot_text(current_plot, label, x_label, value)
        current_plot.setTitle(
            self._text(
                f"电流｜{row.get('Stage', '')}｜{row['Direction']}",
                f"Current | {row.get('Stage', '')} | {row['Direction']}",
            )
        )

        force_plot = self.response_plot_area.addPlot(row=1, col=0)
        force_plot.setXLink(current_plot)
        self._axis_style(force_plot, self._text("阻尼力", "Damping force"), "kN")
        force_kn = data[LOAD].to_numpy(float) / 1000.0
        force_plot.plot(t_s, force_kn, pen=signal_pen)
        force_plot.addLine(x=t0_s, pen=marker_pen)

        force_levels = (
            ("1%", float(row["F1 N"]) / 1000.0),
            ("63%", float(row["F63 N"]) / 1000.0),
            ("90%", float(row["F90 N"]) / 1000.0),
            ("100%", float(row["F100 N"]) / 1000.0),
        )
        for label, value in force_levels:
            force_plot.addLine(y=value, pen=marker_pen)
            self._add_plot_text(force_plot, label, x_label, value)

        response_markers = (
            ("t1", float(row["Dead Time t1 ms"])),
            ("t63", float(row["Switch Time t63 ms"])),
            ("t90", float(row["Switch Time t90 ms"])),
        )
        finite_marker_positions: list[tuple[str, float]] = []
        for label, elapsed_ms in response_markers:
            if np.isfinite(elapsed_ms):
                x_value = t0_s + elapsed_ms / 1000.0
                force_plot.addLine(x=x_value, pen=marker_pen)
                finite_marker_positions.append((label, x_value))
        if len(force_kn):
            y_min = float(np.nanmin(force_kn))
            y_max = float(np.nanmax(force_kn))
            y_span = max(y_max - y_min, 0.1)
            for index, (label, x_value) in enumerate(finite_marker_positions):
                y = y_max - (0.08 + 0.12 * index) * y_span
                self._add_plot_text(
                    force_plot,
                    label,
                    x_value,
                    y,
                    anchor=(0.5, 0.5),
                )

        # Direction (Rebound/Compression) remains available in the event title,
        # selector and result table. Do not repeat "复原（+）/压缩（-）" inside
        # the force plot itself.

        velocity_plot = self.response_plot_area.addPlot(row=2, col=0)
        velocity_plot.setXLink(current_plot)
        self._axis_style(velocity_plot, self._text("速度", "Velocity"), "m/s")
        velocity_plot.plot(t_s, data[VELOCITY].to_numpy(float), pen=signal_pen)
        velocity_plot.addLine(x=t0_s, pen=marker_pen)
        # Target velocity is still used by the response-event gating and stays
        # in the result table, but its horizontal line/text are intentionally
        # omitted from the detail plot.

        current_plot.setXRange(x_left, x_right, padding=0.01)


# Consolidated behavior layer 078
class _Layer078(_Layer077):
    """V0.7.8 response presentation corrections.

    - Chinese UI localizes response stage/direction values instead of exposing
      internal English enum text.
    - The velocity detail plot shows the target-speed horizontal dashed line.
      The old t0 vertical line is removed from the velocity plot.
    - Target-speed text remains omitted to keep the velocity plot uncluttered.
    """

    def _localized_stage(self, value: object) -> str:
        text = str(value or "")
        if self.window.language != "zh_CN":
            return text
        replacements = (
            ("Soft", "软"),
            ("Hard", "硬"),
            ("Medium", "中间"),
            ("Mid", "中间"),
        )
        for source, target in replacements:
            text = text.replace(source, target)
        return text

    def _localized_direction(self, value: object, *, signed: bool = False) -> str:
        text = str(value or "")
        if self.window.language != "zh_CN":
            if not signed:
                return text
            if text == "Rebound":
                return "Rebound (+)"
            if text == "Compression":
                return "Compression (-)"
            return text
        if text == "Rebound":
            return "复原(+)" if signed else "复原"
        if text == "Compression":
            return "压缩(-)" if signed else "压缩"
        return text

    def _rebuild_response_event_combo(self) -> None:
        if self.response_result is None or self.response_result.events.empty:
            return
        current_event = self.response_event_combo.currentData()
        self.response_event_combo.blockSignals(True)
        self.response_event_combo.clear()
        for _, row in self.response_result.events.iterrows():
            stage = self._localized_stage(row.get("Stage", ""))
            direction = self._localized_direction(row.get("Direction", ""), signed=True)
            current_transition = str(row.get("Current Transition", ""))
            target = float(row["Target Velocity m/s"])
            if self.window.language == "zh_CN":
                text = f"{stage} | {direction} | {current_transition} | 目标速度 {target:+.4g} m/s"
            else:
                text = f"{stage} | {direction} | {current_transition} | Target {target:+.4g} m/s"
            self.response_event_combo.addItem(text, int(row["Event ID"]))
        restore_index = self.response_event_combo.findData(current_event)
        self.response_event_combo.setCurrentIndex(restore_index if restore_index >= 0 else 0)
        self.response_event_combo.blockSignals(False)

    def analyze_response(self):
        super().analyze_response()
        if self.response_result is not None and not self.response_result.events.empty:
            self._rebuild_response_event_combo()
            self._fill_table(self.response_table, self.response_result.events)
            self.refresh_response_plot()

    def apply_language(self, language: str):
        super().apply_language(language)
        if getattr(self, "response_result", None) is not None:
            self._rebuild_response_event_combo()
            self._fill_table(self.response_table, self.response_result.events)
            self.refresh_response_plot()

    def _fill_table(self, table, frame):
        super()._fill_table(table, frame)
        if (
            self.window.language != "zh_CN"
            or not hasattr(self, "response_table")
            or table is not self.response_table
        ):
            return
        columns = [str(column) for column in frame.columns]
        for column_name, translator in (
            ("Stage", self._localized_stage),
            ("Direction", self._localized_direction),
        ):
            if column_name not in columns:
                continue
            column_index = columns.index(column_name)
            for row_index in range(table.rowCount()):
                item = table.item(row_index, column_index)
                if item is not None:
                    item.setText(translator(item.text()))

    def refresh_response_plot(self):
        super().refresh_response_plot()
        if self.response_result is None or self.response_result.events.empty:
            return

        event_id = self.response_event_combo.currentData()
        if event_id is None:
            event_id = int(self.response_result.events["Event ID"].iloc[0])
        row = self.response_result.events[
            self.response_result.events["Event ID"] == event_id
        ].iloc[0]

        current_plot = self.response_plot_area.getItem(0, 0)
        velocity_plot = self.response_plot_area.getItem(2, 0)
        if current_plot is not None:
            stage = self._localized_stage(row.get("Stage", ""))
            direction = self._localized_direction(row.get("Direction", ""))
            current_plot.setTitle(
                self._text(
                    f"电流｜{stage}｜{direction}",
                    f"Current | {stage} | {direction}",
                )
            )

        if velocity_plot is not None:
            # V0.7.7 left a t0 vertical marker on the velocity subplot. For the
            # engineering view the useful reference here is target velocity,
            # so remove vertical reference lines and restore one horizontal
            # target-speed dashed line. Keep the target text hidden.
            for item in list(velocity_plot.items):
                if isinstance(item, self.pg.InfiniteLine):
                    velocity_plot.removeItem(item)
            target_velocity = float(row["Target Velocity m/s"])
            velocity_plot.addLine(y=target_velocity, pen=self._marker_pen())


# Consolidated behavior layer 079
class _Layer079(_Layer078):
    """V0.7.9 response-control cleanup.

    The end-average fraction remains an internal engineering parameter fixed at
    2% for the current implementation, but it is no longer exposed in the main
    response toolbar. The project t90 limit stays visible because it drives the
    PASS/FAIL judgement in the result table.
    """

    def _build_response_page(self):
        super()._build_response_page()

        # Keep the existing algorithm parameter for compatibility/export, but
        # remove it from the normal operator interface. Audi uses the specified
        # final-2% mean; BMW currently keeps the same engineering default until
        # its exact F Anfang/F Ende extraction rule is confirmed.
        self.response_end_fraction.setValue(2.0)
        self.response_end_label.hide()
        self.response_end_fraction.hide()

        # Zero means no project acceptance limit. Present that state explicitly
        # instead of the misleading numeric text "0.00 ms".
        self.response_t90_limit.setSpecialValueText(
            self._text("未设置", "Not set")
        )

    def apply_language(self, language: str):
        super().apply_language(language)
        if hasattr(self, "response_end_fraction"):
            self.response_end_fraction.setValue(2.0)
            self.response_end_label.hide()
            self.response_end_fraction.hide()
            self.response_t90_limit.setSpecialValueText(
                self._text("未设置", "Not set")
            )


# Consolidated behavior layer 080
class _Layer080(_Layer079):
    """V0.8.0 response target-speed configuration.

    OEM profile and test speed are deliberately separated. BMW/Audi populate
    convenient default speeds, but the operator may enter any positive customer
    speed list (for example 0.1, 0.3, 0.6, 1.0 m/s).
    """

    def _build_response_page(self):
        super()._build_response_page()

        controls = self.response_page.layout().itemAt(0).layout()
        self.response_target_speed_label = self.QtWidgets.QLabel()
        self.response_target_speeds = self.QtWidgets.QLineEdit()
        self.response_target_speeds.setMinimumWidth(190)
        self.response_target_speeds.setMaximumWidth(260)
        self.response_target_speeds.setClearButtonEnabled(True)

        standard_index = controls.indexOf(self.response_standard)
        insert_at = standard_index + 1 if standard_index >= 0 else 0
        controls.insertWidget(insert_at, self.response_target_speed_label)
        controls.insertWidget(insert_at + 1, self.response_target_speeds)

        self.response_standard.currentIndexChanged.connect(
            self._reset_target_speeds_for_standard
        )
        self._reset_target_speeds_for_standard()
        self._apply_target_speed_text()

    def _current_standard(self) -> ResponseStandard:
        return ResponseStandard(self.response_standard.currentData())

    def _format_speed_list(self, values) -> str:
        return ", ".join(f"{float(value):g}" for value in values)

    def _reset_target_speeds_for_standard(self, *_args):
        if not hasattr(self, "response_target_speeds"):
            return
        self.response_target_speeds.setText(
            self._format_speed_list(default_target_speeds(self._current_standard()))
        )

    def _apply_target_speed_text(self):
        if not hasattr(self, "response_target_speed_label"):
            return
        self.response_target_speed_label.setText(
            self._text("目标速度", "Target speed")
        )
        self.response_target_speeds.setPlaceholderText(
            self._text(
                "例如 0.1, 0.3, 0.6, 1.0",
                "e.g. 0.1, 0.3, 0.6, 1.0",
            )
        )
        self.response_target_speeds.setToolTip(
            self._text(
                "可输入任意客户目标速度，单位 m/s；多个速度用逗号分隔。BMW 默认 0.131, 0.524, 1.048。",
                "Enter any customer target speeds in m/s, separated by commas. BMW defaults: 0.131, 0.524, 1.048.",
            )
        )

    def apply_language(self, language: str):
        super().apply_language(language)
        self._apply_target_speed_text()

    def analyze_response(self):
        try:
            target_speeds = parse_target_speeds(self.response_target_speeds.text())
        except ValueError as exc:
            self.QtWidgets.QMessageBox.information(
                self.window,
                self._text("目标速度", "Target speed"),
                str(exc),
            )
            return

        original_analyzer = _dynamic_gui_v074_module.analyze_response_time_v074

        def configured_analyzer(dataset, config, *, target_speed_tolerance=0.10):
            return analyze_response_time_v080(
                dataset,
                config,
                target_speeds_mps=target_speeds,
                target_speed_tolerance=target_speed_tolerance,
            )

        _dynamic_gui_v074_module.analyze_response_time_v074 = configured_analyzer
        try:
            super().analyze_response()
        finally:
            _dynamic_gui_v074_module.analyze_response_time_v074 = original_analyzer


# Consolidated behavior layer 081
class _Layer081(_Layer080):
    """V0.8.1 engineering notation for response-detail plots.

    Presentation changes only:
    - current reference labels use engineering subscripts: I₁₀%, I₁₀₀%;
    - force reference labels use F with subscripted percentage levels;
    - response-time markers use t with subscripted percentage levels and show
      the calculated result directly beside the marker, e.g. t₉₀% = 7.21 ms;
    - all reference lines remain dashed; measured signals remain solid.
    """

    def _format_response_ms(self, value: float) -> str:
        return f"{value:.2f}"

    def refresh_response_plot(self):
        self.response_plot_area.clear()
        if self.response_result is None or self.response_result.events.empty:
            return

        event_id = self.response_event_combo.currentData()
        if event_id is None:
            event_id = int(self.response_result.events["Event ID"].iloc[0])
        row = self.response_result.events[
            self.response_result.events["Event ID"] == event_id
        ].iloc[0]
        self._sync_response_table_selection(int(event_id))

        start = float(row.get("Display Start s", row["Segment Start s"]))
        end = float(row.get("Display End s", row["Segment End s"]))
        data = self.response_result.processed[
            self.response_result.processed[TIME].between(start, end)
        ]
        if data.empty:
            return

        t_s = data[TIME].to_numpy(float)
        t0_s = float(row["t0 s"])
        x_left = float(t_s[0])
        x_right = float(t_s[-1])
        x_label = x_left + 0.02 * max(x_right - x_left, 1e-9)
        marker_pen = self._marker_pen()
        signal_pen = self._signal_pen()

        stage = self._localized_stage(row.get("Stage", ""))
        direction = self._localized_direction(row.get("Direction", ""))

        current_plot = self.response_plot_area.addPlot(row=0, col=0)
        self._axis_style(current_plot, self._text("阀电流", "Valve current"), "A")
        current_plot.plot(t_s, data[CURRENT].to_numpy(float), pen=signal_pen)
        current_plot.addLine(x=t0_s, pen=marker_pen)
        current_levels = (
            ("I₁₀%", float(row["Trigger Current A"])),
            ("I₁₀₀%", float(row["Current 100% A"])),
        )
        for label, value in current_levels:
            current_plot.addLine(y=value, pen=marker_pen)
            self._add_plot_text(current_plot, label, x_label, value)
        current_plot.setTitle(
            self._text(
                f"电流｜{stage}｜{direction}",
                f"Current | {stage} | {direction}",
            )
        )

        force_plot = self.response_plot_area.addPlot(row=1, col=0)
        force_plot.setXLink(current_plot)
        self._axis_style(force_plot, self._text("阻尼力", "Damping force"), "kN")
        force_kn = data[LOAD].to_numpy(float) / 1000.0
        force_plot.plot(t_s, force_kn, pen=signal_pen)
        force_plot.addLine(x=t0_s, pen=marker_pen)

        force_levels = (
            ("F₁%", float(row["F1 N"]) / 1000.0),
            ("F₆₃%", float(row["F63 N"]) / 1000.0),
            ("F₉₀%", float(row["F90 N"]) / 1000.0),
            ("F₁₀₀%", float(row["F100 N"]) / 1000.0),
        )
        for label, value in force_levels:
            force_plot.addLine(y=value, pen=marker_pen)
            self._add_plot_text(force_plot, label, x_label, value)

        response_markers = (
            ("t₁%", float(row["Dead Time t1 ms"])),
            ("t₆₃%", float(row["Switch Time t63 ms"])),
            ("t₉₀%", float(row["Switch Time t90 ms"])),
        )
        finite_marker_positions: list[tuple[str, float, float]] = []
        for label, elapsed_ms in response_markers:
            if np.isfinite(elapsed_ms):
                x_value = t0_s + elapsed_ms / 1000.0
                force_plot.addLine(x=x_value, pen=marker_pen)
                finite_marker_positions.append((label, x_value, elapsed_ms))

        if len(force_kn):
            y_min = float(np.nanmin(force_kn))
            y_max = float(np.nanmax(force_kn))
            y_span = max(y_max - y_min, 0.1)
            for index, (label, x_value, elapsed_ms) in enumerate(finite_marker_positions):
                y = y_max - (0.08 + 0.12 * index) * y_span
                result_text = f"{label} = {self._format_response_ms(elapsed_ms)} ms"
                self._add_plot_text(
                    force_plot,
                    result_text,
                    x_value,
                    y,
                    anchor=(0.5, 0.5),
                )

        velocity_plot = self.response_plot_area.addPlot(row=2, col=0)
        velocity_plot.setXLink(current_plot)
        self._axis_style(velocity_plot, self._text("速度", "Velocity"), "m/s")
        velocity_plot.plot(t_s, data[VELOCITY].to_numpy(float), pen=signal_pen)
        target_velocity = float(row["Target Velocity m/s"])
        velocity_plot.addLine(y=target_velocity, pen=marker_pen)

        current_plot.setXRange(x_left, x_right, padding=0.01)


# Consolidated behavior layer 083
class _Layer083(_Layer081):
    """V0.8.3 response plots with readable, traceable annotations."""

    def _add_response_text(self, plot, text, x, y, *, anchor=(0, 1), bold=True):
        foreground = getattr(self.window, "_plot_foreground_color", "#202020")
        background = getattr(self.window, "_plot_background_color", "#ffffff")
        # An opaque plot-matched fill keeps dashed reference lines from running
        # through glyphs while remaining correct on every selectable background.
        item = self.pg.TextItem(
            text=text,
            color=foreground,
            anchor=anchor,
            fill=self.pg.mkBrush(background),
        )
        font = self.QtWidgets.QApplication.font()
        font.setPointSize(max(12, font.pointSize()))
        font.setBold(bool(bold))
        item.setFont(font)
        item.setPos(float(x), float(y))
        plot.addItem(item)
        return item

    def refresh_response_plot(self):
        self.response_plot_area.clear()
        if self.response_result is None or self.response_result.events.empty:
            return

        event_id = self.response_event_combo.currentData()
        if event_id is None:
            event_id = int(self.response_result.events["Event ID"].iloc[0])
        row = self.response_result.events[
            self.response_result.events["Event ID"] == event_id
        ].iloc[0]
        self._sync_response_table_selection(int(event_id))

        start = float(row.get("Display Start s", row["Segment Start s"]))
        end = float(row.get("Display End s", row["Segment End s"]))
        # F100 is calculated from the target-speed endpoint window. Keep that
        # window visible so every force reference line crosses measured data.
        end = max(end, float(row.get("Target Window End s", end)))
        data = self.response_result.processed[
            self.response_result.processed[TIME].between(start, end)
        ]
        if data.empty:
            return

        t_s = data[TIME].to_numpy(float)
        t0_s = float(row["t0 s"])
        x_left = float(t_s[0])
        x_right = float(t_s[-1])
        x_span = max(x_right - x_left, 1e-9)
        x_label = x_left + 0.02 * x_span
        marker_pen = self._marker_pen()
        signal_pen = self._signal_pen()

        stage = self._localized_stage(row.get("Stage", ""))
        direction = self._localized_direction(row.get("Direction", ""))

        current_plot = self.response_plot_area.addPlot(row=0, col=0)
        self._axis_style(current_plot, self._text("阀电流", "Valve current"), "A")
        current_plot.plot(t_s, data[CURRENT].to_numpy(float), pen=signal_pen)
        current_plot.addLine(x=t0_s, pen=marker_pen)
        for label, value in (
            ("I₁₀%", float(row["Trigger Current A"])),
            ("I₁₀₀%", float(row["Current 100% A"])),
        ):
            current_plot.addLine(y=value, pen=marker_pen)
            self._add_response_text(current_plot, label, x_label, value)
        current_plot.setTitle(
            self._text(
                f"电流｜{stage}｜{direction}",
                f"Current | {stage} | {direction}",
            )
        )

        force_plot = self.response_plot_area.addPlot(row=1, col=0)
        force_plot.setXLink(current_plot)
        self._axis_style(force_plot, self._text("阻尼力", "Damping force"), "kN")
        force_kn = data[LOAD].to_numpy(float) / 1000.0
        force_plot.plot(t_s, force_kn, pen=signal_pen)
        force_plot.addLine(x=t0_s, pen=marker_pen)

        for label, value in (
            ("F₁%", float(row["F1 N"]) / 1000.0),
            ("F₆₃%", float(row["F63 N"]) / 1000.0),
            ("F₉₀%", float(row["F90 N"]) / 1000.0),
            ("F₁₀₀%", float(row["F100 N"]) / 1000.0),
        ):
            force_plot.addLine(y=value, pen=marker_pen)
            self._add_response_text(force_plot, label, x_label, value)

        finite_marker_positions = []
        for label, elapsed_ms in (
            ("t₁%", float(row["Dead Time t1 ms"])),
            ("t₆₃%", float(row["Switch Time t63 ms"])),
            ("t₉₀%", float(row["Switch Time t90 ms"])),
        ):
            if np.isfinite(elapsed_ms):
                x_value = t0_s + elapsed_ms / 1000.0
                force_plot.addLine(x=x_value, pen=marker_pen)
                finite_marker_positions.append((label, x_value, elapsed_ms))

        if len(force_kn):
            y_min = float(np.nanmin(force_kn))
            y_max = float(np.nanmax(force_kn))
            y_span = max(y_max - y_min, 0.1)
            x_offset = 0.015 * x_span
            for index, (label, x_value, elapsed_ms) in enumerate(finite_marker_positions):
                y = y_max - (0.06 + 0.16 * index) * y_span
                if x_value <= x_right - 0.22 * x_span:
                    text_x, anchor = x_value + x_offset, (0, 1)
                else:
                    text_x, anchor = x_value - x_offset, (1, 1)
                self._add_response_text(
                    force_plot,
                    f"{label} = {self._format_response_ms(elapsed_ms)} ms",
                    text_x,
                    y,
                    anchor=anchor,
                )

        velocity_plot = self.response_plot_area.addPlot(row=2, col=0)
        velocity_plot.setXLink(current_plot)
        self._axis_style(velocity_plot, self._text("速度", "Velocity"), "m/s")
        velocity_plot.plot(t_s, data[VELOCITY].to_numpy(float), pen=signal_pen)
        velocity_plot.addLine(y=float(row["Target Velocity m/s"]), pen=marker_pen)

        current_plot.setXRange(x_left, x_right, padding=0.01)


# Consolidated behavior layer 084
class _Layer084(_Layer083):
    """V0.8.4 response plots with transparent labels and curve-intersection markers."""

    def _axis_title_font(self, plot):
        """Return a normal-weight font matching the plot's left-axis title font."""
        axis = plot.getAxis("left")
        label_item = getattr(axis, "label", None)
        font = label_item.font() if label_item is not None else self.QtWidgets.QApplication.font()
        font.setBold(False)
        return font

    def _add_response_text(self, plot, text, x, y, *, anchor=(0, 1), bold=False):
        foreground = getattr(self.window, "_plot_foreground_color", "#202020")
        # Deliberately omit fill/background: engineering annotations remain transparent.
        item = self.pg.TextItem(
            text=text,
            color=foreground,
            anchor=anchor,
            fill=None,
        )
        font = self._axis_title_font(plot)
        font.setBold(False)
        item.setFont(font)
        item.setPos(float(x), float(y))
        item.setZValue(30)
        plot.addItem(item)
        return item

    @staticmethod
    def _curve_y_at(x_values, y_values, x_value: float) -> float:
        x = np.asarray(x_values, dtype=float)
        y = np.asarray(y_values, dtype=float)
        finite = np.isfinite(x) & np.isfinite(y)
        if finite.sum() < 2:
            return float("nan")
        x = x[finite]
        y = y[finite]
        order = np.argsort(x, kind="stable")
        x = x[order]
        y = y[order]
        if x_value < x[0] or x_value > x[-1]:
            return float("nan")
        return float(np.interp(float(x_value), x, y))

    def _add_curve_intersection_marker(self, plot, x_value: float, y_value: float):
        if not (np.isfinite(x_value) and np.isfinite(y_value)):
            return None
        foreground = getattr(self.window, "_plot_foreground_color", "#202020")
        marker = self.pg.ScatterPlotItem(
            x=[float(x_value)],
            y=[float(y_value)],
            symbol="o",
            size=8,
            pen=self.pg.mkPen(foreground, width=1.2),
            brush=self.pg.mkBrush(foreground),
        )
        marker.setZValue(25)
        plot.addItem(marker)
        return marker

    @staticmethod
    def _vertical_text_position(y_value: float, y_min: float, y_max: float, index: int = 0):
        y_span = max(float(y_max) - float(y_min), 1e-9)
        offset = max(0.055 * y_span, 1e-6)
        near_top = y_value >= y_max - 0.15 * y_span
        near_bottom = y_value <= y_min + 0.15 * y_span
        if near_top:
            above = False
        elif near_bottom:
            above = True
        else:
            above = index % 2 == 0
        if above:
            return float(y_value + offset), (0.5, 1.0)
        return float(y_value - offset), (0.5, 0.0)

    def refresh_response_plot(self):
        self.response_plot_area.clear()
        if self.response_result is None or self.response_result.events.empty:
            return

        event_id = self.response_event_combo.currentData()
        if event_id is None:
            event_id = int(self.response_result.events["Event ID"].iloc[0])
        row = self.response_result.events[
            self.response_result.events["Event ID"] == event_id
        ].iloc[0]
        self._sync_response_table_selection(int(event_id))

        start = float(row.get("Display Start s", row["Segment Start s"]))
        end = float(row.get("Display End s", row["Segment End s"]))
        # Keep the complete target-speed endpoint window visible so F100 remains traceable.
        end = max(end, float(row.get("Target Window End s", end)))
        data = self.response_result.processed[
            self.response_result.processed[TIME].between(start, end)
        ]
        if data.empty:
            return

        t_s = data[TIME].to_numpy(float)
        current_a = data[CURRENT].to_numpy(float)
        force_kn = data[LOAD].to_numpy(float) / 1000.0
        velocity = data[VELOCITY].to_numpy(float)
        t0_s = float(row["t0 s"])
        x_left = float(t_s[0])
        x_right = float(t_s[-1])
        x_span = max(x_right - x_left, 1e-9)
        x_label = x_left + 0.02 * x_span
        marker_pen = self._marker_pen()
        signal_pen = self._signal_pen()

        stage = self._localized_stage(row.get("Stage", ""))
        direction = self._localized_direction(row.get("Direction", ""))

        current_plot = self.response_plot_area.addPlot(row=0, col=0)
        self._axis_style(current_plot, self._text("阀电流", "Valve current"), "A")
        current_plot.plot(t_s, current_a, pen=signal_pen)
        current_plot.addLine(x=t0_s, pen=marker_pen)
        for label, value in (
            ("I₁₀%", float(row["Trigger Current A"])),
            ("I₁₀₀%", float(row["Current 100% A"])),
        ):
            current_plot.addLine(y=value, pen=marker_pen)
            self._add_response_text(current_plot, label, x_label, value, anchor=(0, 1))

        current_t0_y = self._curve_y_at(t_s, current_a, t0_s)
        self._add_curve_intersection_marker(current_plot, t0_s, current_t0_y)
        if np.isfinite(current_t0_y):
            current_min = float(np.nanmin(current_a))
            current_max = float(np.nanmax(current_a))
            text_y, anchor = self._vertical_text_position(current_t0_y, current_min, current_max, 0)
            self._add_response_text(current_plot, "t₀", t0_s, text_y, anchor=anchor)

        current_plot.setTitle(
            self._text(
                f"电流｜{stage}｜{direction}",
                f"Current | {stage} | {direction}",
            )
        )

        force_plot = self.response_plot_area.addPlot(row=1, col=0)
        force_plot.setXLink(current_plot)
        self._axis_style(force_plot, self._text("阻尼力", "Damping force"), "kN")
        force_plot.plot(t_s, force_kn, pen=signal_pen)
        force_plot.addLine(x=t0_s, pen=marker_pen)

        for label, value in (
            ("F₁%", float(row["F1 N"]) / 1000.0),
            ("F₆₃%", float(row["F63 N"]) / 1000.0),
            ("F₉₀%", float(row["F90 N"]) / 1000.0),
            ("F₁₀₀%", float(row["F100 N"]) / 1000.0),
        ):
            force_plot.addLine(y=value, pen=marker_pen)
            self._add_response_text(force_plot, label, x_label, value, anchor=(0, 1))

        vertical_markers = [("t₀", t0_s, None)]
        for label, elapsed_ms in (
            ("t₁%", float(row["Dead Time t1 ms"])),
            ("t₆₃%", float(row["Switch Time t63 ms"])),
            ("t₉₀%", float(row["Switch Time t90 ms"])),
        ):
            if np.isfinite(elapsed_ms):
                x_value = t0_s + elapsed_ms / 1000.0
                force_plot.addLine(x=x_value, pen=marker_pen)
                vertical_markers.append((label, x_value, elapsed_ms))

        if len(force_kn):
            y_min = float(np.nanmin(force_kn))
            y_max = float(np.nanmax(force_kn))
            for index, (label, x_value, elapsed_ms) in enumerate(vertical_markers):
                y_value = self._curve_y_at(t_s, force_kn, x_value)
                self._add_curve_intersection_marker(force_plot, x_value, y_value)
                if not np.isfinite(y_value):
                    continue
                text_y, anchor = self._vertical_text_position(y_value, y_min, y_max, index)
                text = label if elapsed_ms is None else f"{label} = {self._format_response_ms(elapsed_ms)} ms"
                # The label shares the marker x-coordinate and moves only vertically.
                self._add_response_text(force_plot, text, x_value, text_y, anchor=anchor)

        velocity_plot = self.response_plot_area.addPlot(row=2, col=0)
        velocity_plot.setXLink(current_plot)
        self._axis_style(velocity_plot, self._text("速度", "Velocity"), "m/s")
        velocity_plot.plot(t_s, velocity, pen=signal_pen)
        velocity_plot.addLine(y=float(row["Target Velocity m/s"]), pen=marker_pen)

        current_plot.setXRange(x_left, x_right, padding=0.01)


class DynamicPagesController(_Layer084):
    """Canonical V0.9 runtime controller preserving V0.8.4 behavior.

    Production code no longer imports dynamic_gui_v07x/v08x modules.
    Historical files remain only as regression references during the freeze.
    """
    pass
    