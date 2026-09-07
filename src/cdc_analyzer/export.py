from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from .analysis import AnalysisResult
from .formatting import excel_number_format
from .quality import build_block_quality


RESULT_SHEETS = {"Summary", "Run Detail", "Cycle Detail", "Data Quality", "Processed Data"}


def _autosize_sheet(ws) -> None:
    ws.freeze_panes = "A2"
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for col_idx, column_cells in enumerate(ws.columns, start=1):
        max_len = 0
        for cell in column_cells:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 10), 36)


def _apply_number_formats(ws) -> None:
    """Apply presentation precision without rounding the stored numeric values."""
    if ws.title not in RESULT_SHEETS or ws.max_row < 2:
        return
    for col_idx in range(1, ws.max_column + 1):
        header = ws.cell(row=1, column=col_idx).value
        if not isinstance(header, str):
            continue
        number_format = excel_number_format(header)
        for row_idx in range(2, ws.max_row + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                cell.number_format = number_format


def export_xlsx(result: AnalysisResult, path: str | Path, include_raw: bool = False) -> Path:
    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        path = path.with_suffix(".xlsx")
    path.parent.mkdir(parents=True, exist_ok=True)

    settings_df = pd.DataFrame([{"Setting": k, "Value": v} for k, v in result.settings.items()])
    quality_df = build_block_quality(result.processed)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        result.summary.to_excel(writer, sheet_name="Summary", index=False)
        result.runs.to_excel(writer, sheet_name="Run Detail", index=False)
        result.cycles.to_excel(writer, sheet_name="Cycle Detail", index=False)
        quality_df.to_excel(writer, sheet_name="Data Quality", index=False)
        settings_df.to_excel(writer, sheet_name="Settings", index=False)
        if include_raw:
            result.processed.to_excel(writer, sheet_name="Processed Data", index=False)

    wb = load_workbook(path)
    for ws in wb.worksheets:
        _apply_number_formats(ws)
        _autosize_sheet(ws)
    wb.save(path)
    return path
