from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from .dynamic_analysis import HysteresisAnalysisResult, ResponseAnalysisResult
from .formatting import excel_number_format


def _autosize(ws) -> None:
    ws.freeze_panes = "A2"
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for index, cells in enumerate(ws.columns, start=1):
        width = 10
        for cell in cells:
            if cell.value is not None:
                width = max(width, min(len(str(cell.value)) + 2, 42))
        ws.column_dimensions[get_column_letter(index)].width = width


def _format_numbers(ws) -> None:
    if ws.max_row < 2:
        return
    for col in range(1, ws.max_column + 1):
        header = ws.cell(1, col).value
        if not isinstance(header, str):
            continue
        fmt = excel_number_format(header)
        for row in range(2, ws.max_row + 1):
            cell = ws.cell(row, col)
            if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                cell.number_format = fmt


def _settings_frame(settings: dict[str, object]) -> pd.DataFrame:
    return pd.DataFrame([{"Setting": key, "Value": value} for key, value in settings.items()])


def export_response_xlsx(
    result: ResponseAnalysisResult,
    path: str | Path,
    include_processed: bool = True,
) -> Path:
    out = Path(path).with_suffix(".xlsx")
    out.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        result.events.to_excel(writer, sheet_name="Response Summary", index=False)
        _settings_frame(result.settings).to_excel(writer, sheet_name="Settings", index=False)
        if include_processed:
            result.processed.to_excel(writer, sheet_name="Processed Data", index=False)
    workbook = load_workbook(out)
    for sheet in workbook.worksheets:
        _format_numbers(sheet)
        _autosize(sheet)
    workbook.save(out)
    return out


def export_hysteresis_xlsx(
    result: HysteresisAnalysisResult,
    path: str | Path,
    include_processed: bool = False,
) -> Path:
    out = Path(path).with_suffix(".xlsx")
    out.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        result.summary.to_excel(writer, sheet_name="Hysteresis Summary", index=False)
        result.runs.to_excel(writer, sheet_name="Run Detail", index=False)
        _settings_frame(result.settings).to_excel(writer, sheet_name="Settings", index=False)
        if include_processed:
            result.processed.to_excel(writer, sheet_name="Processed Data", index=False)
    workbook = load_workbook(out)
    for sheet in workbook.worksheets:
        _format_numbers(sheet)
        _autosize(sheet)
    workbook.save(out)
    return out
