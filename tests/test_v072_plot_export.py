from __future__ import annotations

import os
from pathlib import Path

import pandas as pd
import pytest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
pytest.importorskip("PySide6")
pytest.importorskip("pyqtgraph")
PIL = pytest.importorskip("PIL.Image")


def test_nice_step_uses_engineering_intervals():
    from cdc_analyzer.gui_release_v072 import _nice_step

    assert _nice_step(0.036) == pytest.approx(0.01)
    assert _nice_step(0.11) == pytest.approx(0.02)
    assert _nice_step(1.8) == pytest.approx(0.5)
    assert _nice_step(6200.0) == pytest.approx(2000.0)


def test_v072_marker_dash_and_high_resolution_png(tmp_path: Path):
    from PySide6 import QtCore, QtWidgets
    from cdc_analyzer.gui_release_v072 import DASH_PATTERN, _build_release_gui_classes

    app = QtWidgets.QApplication.instance() or QtWidgets.QApplication([])
    MainWindow = _build_release_gui_classes()
    window = MainWindow()

    pen = window._marker_pen_v07()
    assert pen.style() == QtCore.Qt.PenStyle.CustomDashLine
    assert list(pen.dashPattern()) == pytest.approx(DASH_PATTERN)

    widget = window.dynamic_pages.response_plot_area
    widget.resize(420, 240)
    app.processEvents()
    output = tmp_path / "hires.png"
    window._save_widget_highres_v072(widget, output, scale=2.0)
    with PIL.open(output) as image:
        assert image.width == widget.width() * 2
        assert image.height == widget.height() * 2

    window.close()
    app.processEvents()


def test_response_excel_embeds_overview_and_all_stage_figures(tmp_path: Path):
    from openpyxl import load_workbook
    from PIL import Image
    from cdc_analyzer.dynamic_analysis import ResponseAnalysisResult
    from cdc_analyzer.dynamic_export import export_response_xlsx

    image_paths = []
    for index in range(3):
        path = tmp_path / f"figure_{index}.png"
        Image.new("RGB", (900, 480), "white").save(path)
        image_paths.append(path)

    result = ResponseAnalysisResult(
        processed=pd.DataFrame({"Running Time": [0.0, 0.1], "Axial Load": [0.0, 100.0]}),
        events=pd.DataFrame(
            {
                "Event ID": [1, 2],
                "Stage": ["Soft → Hard", "Hard → Soft"],
                "Direction": ["Rebound", "Compression"],
                "Switch Time t90 ms": [7.2, 8.1],
            }
        ),
        settings={"OEM Profile": "bmw"},
        source_path=tmp_path / "source.dat",
    )
    out = export_response_xlsx(
        result,
        tmp_path / "response.xlsx",
        stage_figures=[("Stage 1", image_paths[0]), ("Stage 2", image_paths[1])],
        overview_figure=("Overview", image_paths[2]),
        language="zh_CN",
    )
    workbook = load_workbook(out)
    assert "响应汇总" in workbook.sheetnames
    assert "阶段图形" in workbook.sheetnames
    assert "全流程总览" in workbook.sheetnames
    assert len(workbook["阶段图形"]._images) == 2
    assert len(workbook["全流程总览"]._images) == 1
