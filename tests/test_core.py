from __future__ import annotations

from math import pi
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from cdc_analyzer.analysis import AnalyzerConfig, CDCAnalyzer, EvaluationProfile, gas_force_from_pressure
from cdc_analyzer.export import export_xlsx
from cdc_analyzer.formatting import decimals_for_column, format_value
from cdc_analyzer.parser import DataSet, load_test_data


def synthetic_dataset(cycles: int = 2, current: float = 0.8) -> DataSet:
    # Start/end at +stroke so complete max-to-max cycles are explicit.
    n = cycles * 400 + 1
    theta = np.linspace(0, cycles * 2 * pi, n)
    t = theta / (2 * pi) * 4.0
    x = 50.0 * np.cos(theta)
    v_sign = np.sign(-np.sin(theta))  # derivative of cos
    # Rebound positive, compression negative, with a small center-position slope.
    force = np.where(v_sign >= 0, 1000.0 + 2.0 * x, -600.0 + 1.0 * x)
    df = pd.DataFrame(
        {
            "Running Time": t,
            "Axial Displacement": x,
            "Axial Load": force,
            "CDC 1 Current FB_1": current + 0.0002 * np.sin(theta),
            "Block ID": 1,
            "Source Row": np.arange(1, n + 1),
        }
    )
    return DataSet(df, Path("synthetic.dat"), "synthetic")


def test_gas_pressure_formula():
    assert gas_force_from_pressure(0.8, 18.0) == pytest_approx(0.8 * pi * 18.0**2 / 4.0, 1e-9)


def pytest_approx(value: float, tol: float):
    # Tiny helper avoids importing pytest into runtime package code paths.
    import pytest
    return pytest.approx(value, abs=tol)


def test_audi_profile_uses_last_complete_cycle_and_directional_peak():
    result = CDCAnalyzer(AnalyzerConfig(profile=EvaluationProfile.AUDI)).analyze(synthetic_dataset(cycles=2))
    run = result.runs.iloc[0]
    assert int(run["Complete Cycle Count"]) == 2
    # Audi window is center +/- 5 mm for a 100 mm stroke. Peak values occur at window edges.
    assert run["Rebound N"] == pytest_approx(1010.0, 1.0)
    assert run["Compression N"] == pytest_approx(-605.0, 1.0)


def test_window_mean_and_gas_correction():
    cfg = AnalyzerConfig(
        profile=EvaluationProfile.WINDOW_MEAN,
        window_percent=2.0,
        window_basis="amplitude",
        gas_mode="direct",
        gas_force_n=200.0,
        force_channel="corrected",
    )
    result = CDCAnalyzer(cfg).analyze(synthetic_dataset(cycles=1))
    run = result.runs.iloc[0]
    assert run["Rebound N"] == pytest_approx(800.0, 2.0)
    assert run["Compression N"] == pytest_approx(-800.0, 2.0)
    assert result.settings["gas_operation"] == "subtract"
    assert result.settings["gas_correction_n"] == pytest_approx(-200.0, 1e-9)
    assert (result.processed["Gas Force Correction"] == -200.0).all()


def test_gas_correction_can_add_without_overwriting_measured_load():
    cfg = AnalyzerConfig(
        profile=EvaluationProfile.WINDOW_MEAN,
        window_percent=2.0,
        gas_mode="direct",
        gas_operation="add",
        gas_force_n=200.0,
        force_channel="corrected",
    )
    dataset = synthetic_dataset(cycles=1)
    original_load = dataset.data["Axial Load"].copy()
    result = CDCAnalyzer(cfg).analyze(dataset)
    run = result.runs.iloc[0]

    assert run["Rebound N"] == pytest_approx(1200.0, 2.0)
    assert run["Compression N"] == pytest_approx(-400.0, 2.0)
    assert result.settings["gas_operation"] == "add"
    assert result.settings["gas_correction_n"] == pytest_approx(200.0, 1e-9)
    assert (result.processed["Gas Force Correction"] == 200.0).all()
    assert result.processed["Axial Load"].equals(original_load)


def test_pressure_derived_gas_force_uses_selected_operation():
    cfg = AnalyzerConfig(
        gas_mode="pressure",
        gas_operation="add",
        gas_gauge_pressure_mpa=0.8,
        piston_rod_diameter_mm=18.0,
    )
    result = CDCAnalyzer(cfg).analyze(synthetic_dataset(cycles=1))
    expected = gas_force_from_pressure(0.8, 18.0)

    assert result.settings["gas_force_n"] == pytest_approx(expected, 1e-9)
    assert result.settings["gas_correction_n"] == pytest_approx(expected, 1e-9)


def test_cli_accepts_gas_force_operation():
    from cdc_analyzer.cli import build_parser

    parser = build_parser()
    assert parser.parse_args(["sample.dat"]).gas_operation == "subtract"
    assert parser.parse_args(["sample.dat", "--gas-operation", "add"]).gas_operation == "add"


def test_zero_crossing_interpolation():
    cfg = AnalyzerConfig(profile=EvaluationProfile.ZERO_CROSSING)
    result = CDCAnalyzer(cfg).analyze(synthetic_dataset(cycles=1))
    run = result.runs.iloc[0]
    assert run["Rebound N"] == pytest_approx(1000.0, 1.0)
    assert run["Compression N"] == pytest_approx(-600.0, 1.0)


def test_mts_parser_repeated_blocks(tmp_path: Path):
    content = """MTS793|MPT|ENU|1|2|.|/|:|1|0|0|A

Data Acquisition\t\tTime:\t4\ts\nRunning Time\tAxial Displacement\tAxial Load\tCDC 1 Current FB_1\ns\tmm\tN\tA\n0.0\t50\t-500\t0.3\n0.1\t0\t-550\t0.3\n0.2\t-50\t-400\t0.3\n
Data Acquisition\t\tTime:\t8\ts\nRunning Time\tAxial Displacement\tAxial Load\tCDC 1 Current FB_1\ns\tmm\tN\tA\n0.3\t-50\t300\t0.5\n0.4\t0\t500\t0.5\n0.5\t50\t600\t0.5\n"""
    p = tmp_path / "sample.dat"
    p.write_text(content, encoding="utf-8")
    ds = load_test_data(p)
    assert len(ds.data) == 6
    assert ds.metadata["block_count"] == 2
    assert list(ds.data["Block ID"].unique()) == [1, 2]


def test_output_precision_rules():
    assert decimals_for_column("Rebound N") == 0
    assert decimals_for_column("Compression N") == 0
    assert decimals_for_column("Force N") == 0
    assert decimals_for_column("Abs Force N") == 0
    assert decimals_for_column("Reference Damping Force N") == 0
    assert decimals_for_column("Hysteresis N") == 0
    assert decimals_for_column("Current Label A") == 1
    assert decimals_for_column("Current Actual A") == 2
    assert decimals_for_column("Center Position mm") == 2
    assert format_value("Rebound N", 1000.6) == "1001"
    assert format_value("Compression N", -600.6) == "-601"
    assert format_value("Force N", -568.6) == "-569"
    assert format_value("Current Label A", 0.84) == "0.8"
    assert format_value("Current Actual A", 0.836) == "0.84"


def test_excel_export_number_formats(tmp_path: Path):
    result = CDCAnalyzer(AnalyzerConfig(profile=EvaluationProfile.AUDI)).analyze(synthetic_dataset(cycles=1, current=0.84))
    out = export_xlsx(result, tmp_path / "result.xlsx")
    wb = load_workbook(out, data_only=False)
    ws = wb["Summary"]
    headers = {cell.value: cell.column for cell in ws[1]}
    assert ws.cell(2, headers["Rebound N"]).number_format == "0"
    assert ws.cell(2, headers["Compression N"]).number_format == "0"
    assert ws.cell(2, headers["Current Label A"]).number_format == "0.0"
    assert ws.cell(2, headers["Rebound SD N"]).number_format == "0.00"
    # Stored values retain calculation precision; only presentation is rounded.
    assert isinstance(ws.cell(2, headers["Rebound N"]).value, (int, float))
