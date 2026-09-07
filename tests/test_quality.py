from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from cdc_analyzer.analysis import AnalyzerConfig, CDCAnalyzer
from cdc_analyzer.export import export_xlsx
from cdc_analyzer.parser import DataSet
from cdc_analyzer.quality import build_block_quality, overall_quality_status


def frame(times: list[float]) -> pd.DataFrame:
    n = len(times)
    return pd.DataFrame({
        "Running Time": times,
        "Axial Displacement": np.linspace(50.0, -50.0, n),
        "Axial Load": np.linspace(-500.0, 500.0, n),
        "CDC 1 Current FB_1": np.full(n, 0.8),
        "Block ID": np.ones(n, dtype=int),
        "Source Row": np.arange(1, n + 1),
    })


def test_quality_reports_sampling_metrics_without_customer_limit_judgement():
    q = build_block_quality(frame([0.00, 0.02, 0.04, 0.06, 0.08, 0.10]))
    row = q.iloc[0]
    assert row["Status"] == "OK"
    assert abs(float(row["Sample Rate Hz"]) - 50.0) < 1e-9
    assert int(row["Sample Count"]) == 6
    assert overall_quality_status(q) == "OK"


def test_quality_marks_non_increasing_time_invalid():
    q = build_block_quality(frame([0.00, 0.02, 0.04, 0.04, 0.06, 0.08]))
    assert q.iloc[0]["Status"] == "Invalid"
    assert "non-increasing time" in q.iloc[0]["Issues"]
    assert overall_quality_status(q) == "Invalid"


def test_quality_marks_nonfinite_required_channel_invalid():
    f = frame([0.00, 0.02, 0.04, 0.06, 0.08, 0.10])
    f.loc[3, "Axial Load"] = np.nan
    q = build_block_quality(f)
    assert q.iloc[0]["Status"] == "Invalid"
    assert "non-finite samples: 1" in q.iloc[0]["Issues"]


def test_export_contains_quality_and_sweep_sheets(tmp_path: Path):
    # Use a full synthetic cycle so the analyzer can produce normal result tables.
    theta = np.linspace(0, 2 * np.pi, 401)
    t = theta / (2 * np.pi) * 4.0
    x = 50.0 * np.cos(theta)
    force = np.where(np.sign(-np.sin(theta)) >= 0, 1000.0 + 2.0 * x, -600.0 + x)
    raw = pd.DataFrame({
        "Running Time": t,
        "Axial Displacement": x,
        "Axial Load": force,
        "CDC 1 Current FB_1": 0.8 + 0.0002 * np.sin(theta),
        "Block ID": 1,
        "Source Row": np.arange(1, len(theta) + 1),
    })
    result = CDCAnalyzer(AnalyzerConfig()).analyze(DataSet(raw, Path("synthetic.dat"), "synthetic"))
    out = export_xlsx(result, tmp_path / "quality.xlsx")
    wb = load_workbook(out, read_only=True)
    assert "Data Quality" in wb.sheetnames
    assert "Sweep Comparison" in wb.sheetnames
    ws = wb["Data Quality"]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    assert "Sample Rate Hz" in headers
    assert "Status" in headers
